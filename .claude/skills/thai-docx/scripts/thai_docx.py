#!/usr/bin/env python3
"""
thai_docx.py — สร้าง Word document ที่ภาษาไทยเขียนเต็มบรรทัด
โดยแทรก Zero-Width Space (U+200B) ระหว่างคำไทย
เพื่อให้ Word รู้ว่าตัดบรรทัดได้ตรงไหนบ้าง
"""

import re
import sys
import argparse
from pythainlp import word_tokenize
from docx import Document
from docx.shared import Pt, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH


ZWS = "\u200b"  # Zero-Width Space


def insert_zwsp(text: str, engine: str = "newmm") -> str:
    """
    แทรก Zero-Width Space ระหว่างคำไทย
    ข้อความภาษาอังกฤษ, ตัวเลข, URL จะไม่ถูกแตะ
    """
    # แยกส่วนภาษาไทยออกจากส่วนอื่น
    # Pattern: Thai characters + Thai marks
    thai_pattern = re.compile(r'([\u0E00-\u0E7F]+)')

    parts = thai_pattern.split(text)
    result = []

    for part in parts:
        if thai_pattern.fullmatch(part):
            # ตัดคำภาษาไทยแล้วแทรก ZWS
            tokens = word_tokenize(part, engine=engine)
            result.append(ZWS.join(tokens))
        else:
            # ไม่ใช่ภาษาไทย — คงไว้เหมือนเดิม
            result.append(part)

    return "".join(result)


def create_docx(
    paragraphs: list[dict],
    output_path: str,
    font_name: str = "TH Sarabun New",
    font_size: int = 14,
    page_size: str = "A4",
    margins: dict = None,
    line_spacing: float = 1.5,
):
    """
    สร้าง Word document จากรายการ paragraph
    
    paragraphs: list ของ dict แต่ละตัวมี:
        - text: ข้อความ
        - type: "body" | "heading1" | "heading2" | "heading3"
    """
    doc = Document()

    # ตั้งค่า default font
    style = doc.styles["Normal"]
    style.font.name = font_name
    style.font.size = Pt(font_size)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # ตั้งค่า heading styles
    for level, size in [(1, 18), (2, 16), (3, 15)]:
        hstyle = doc.styles[f"Heading {level}"]
        hstyle.font.name = font_name
        hstyle.font.size = Pt(size)
        hstyle.font.bold = True

    # ตั้งค่าหน้ากระดาษ
    section = doc.sections[0]
    if page_size == "A4":
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
    elif page_size == "Letter":
        section.page_width = Cm(21.59)
        section.page_height = Cm(27.94)

    if margins:
        section.top_margin = Cm(margins.get("top", 2.54))
        section.bottom_margin = Cm(margins.get("bottom", 2.54))
        section.left_margin = Cm(margins.get("left", 2.54))
        section.right_margin = Cm(margins.get("right", 2.54))
    else:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    # สร้าง content
    for item in paragraphs:
        text = item["text"]
        ptype = item.get("type", "body")

        # แทรก ZWS ในข้อความภาษาไทย
        processed_text = insert_zwsp(text)

        if ptype == "body":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = line_spacing
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(processed_text)
            run.font.name = font_name
            run.font.size = Pt(font_size)

        elif ptype.startswith("heading"):
            level_str = ptype.replace("heading", "")
            level = int(level_str) if level_str.isdigit() else 1
            p = doc.add_heading(processed_text, level=level)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = font_name

    doc.save(output_path)
    return output_path


# --- CLI ---
if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="สร้าง Word doc ที่ภาษาไทยเขียนเต็มบรรทัด")
    parser.add_argument("input", help="ไฟล์ข้อความ (.txt)")
    parser.add_argument("-o", "--output", default="output.docx", help="ไฟล์ output (.docx)")
    parser.add_argument("--font", default="TH Sarabun New", help="ฟอนต์")
    parser.add_argument("--size", type=int, default=14, help="ขนาดฟอนต์ (pt)")
    args = parser.parse_args()

    with open(args.input, "r", encoding="utf-8") as f:
        raw = f.read()

    # แปลงข้อความเป็น paragraphs (แยกด้วยบรรทัดว่าง)
    blocks = re.split(r"\n{2,}", raw.strip())
    paragraphs = []
    for block in blocks:
        block = block.strip()
        if not block:
            continue
        # ตรวจว่าเป็น heading หรือไม่
        if re.match(r"^\d+(\.\d+)*\s+\S", block) and len(block) < 100:
            depth = block.count(".") 
            level = min(depth + 1, 3)
            paragraphs.append({"text": block, "type": f"heading{level}"})
        else:
            paragraphs.append({"text": block, "type": "body"})

    create_docx(paragraphs, args.output, font_name=args.font, font_size=args.size)
    print(f"✅ Created {args.output}")
